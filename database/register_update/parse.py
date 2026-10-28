"""Parse an FP-ENG-05-01 register into plain Python data (same rules as
database/generate_seed_data.py)."""
import datetime, sys, json
import openpyxl

def d(v):
    if isinstance(v, datetime.datetime): return v.date().isoformat()
    if isinstance(v, datetime.date): return v.isoformat()
    if isinstance(v, str) and v.strip():
        try: return datetime.datetime.strptime(v.strip(), "%d/%m/%Y").date().isoformat()
        except ValueError: return v.strip()
    return None

def s(v):
    if v is None: return None
    t = str(v).strip()
    return t or None

def parse(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb["FP-ENG-05-01 rev.02"] if "FP-ENG-05-01 rev.02" in wb.sheetnames else wb.worksheets[0]
    out, cur = [], None
    for r in range(5, ws.max_row + 1):
        row = [ws.cell(row=r, column=c).value for c in range(1, 17)]
        no, name, model, serial, purpose, rng, acc, oprng, accept, resp, freq, last, nxt, loc, status, remarks = row
        if all(v is None for v in [no, name, model, serial, purpose]): continue
        if no is not None:
            cur = dict(no=s(no), name=s(name), model=s(model), serial=s(serial), location=s(loc), status=s(status), remarks=s(remarks), caps=[], cals=[], row=r)
            out.append(cur)
        if cur is None: continue
        if purpose is not None:
            cur["caps"].append([s(purpose), s(rng), s(acc), s(oprng), s(accept)])
        if last is not None or nxt is not None:
            c = [s(resp), s(freq), d(last), d(nxt)]
            if c not in cur["cals"]: cur["cals"].append(c)
    return wb.sheetnames, out

if __name__ == "__main__":
    names, eq = parse(sys.argv[1])
    print("sheets:", names)
    print("equipment:", len(eq), "caps:", sum(len(e["caps"]) for e in eq), "cals:", sum(len(e["cals"]) for e in eq))
    dates = [c[2] for e in eq for c in e["cals"] if c[2] and c[2][:2] == "20"]
    print("latest calibration date:", max(dates) if dates else None, "| latest next-due:", max(c[3] for e in eq for c in e["cals"] if c[3] and c[3][:2]=="20"))
    json.dump(eq, open(sys.argv[2], "w"), ensure_ascii=False, indent=1)

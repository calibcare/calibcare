"""
Parses 00 FP-ENG-05-01 Rev.02 (equipment calibration register) and emits
database/seed_data.sql -- INSERT statements for equipment,
equipment_capability, and calibration_history matching database/schema.sql.

Run: python3 generate_seed_data.py
"""
import datetime
import openpyxl

SOURCE_XLSX = "../00 FP-ENG-05-01 Rev.02 ทะเบียนควบคุมและประวัติการสอบเที.xlsx"
OUTPUT_SQL = "seed_data.sql"


def esc(value):
    if value is None:
        return "NULL"
    if isinstance(value, datetime.datetime):
        return f"'{value.date().isoformat()}'"
    if isinstance(value, datetime.date):
        return f"'{value.isoformat()}'"
    text = str(value).strip()
    if text == "":
        return "NULL"
    text = text.replace("\\", "\\\\").replace("'", "\\'")
    return f"'{text}'"


def esc_date(value):
    """Same as esc(), but also parses stray text dates like '21/5/2025' (D/M/YYYY)."""
    if isinstance(value, str) and value.strip():
        try:
            parsed = datetime.datetime.strptime(value.strip(), "%d/%m/%Y")
            return f"'{parsed.date().isoformat()}'"
        except ValueError:
            pass
    return esc(value)


def normalize_status(raw):
    if raw is None:
        return "active", None
    text = str(raw).strip()
    if text == "":
        return "active", None
    if text.startswith("ปกติ"):
        return "active", None
    if text.startswith("ยกเลิกการใช้งาน"):
        return "decommissioned", text
    if text.startswith("ส่งแคล"):
        return "sent_for_calibration", text
    return "active", text


def main():
    wb = openpyxl.load_workbook(SOURCE_XLSX, data_only=True)
    ws = wb["FP-ENG-05-01 rev.02"]

    equipment_inserts = []
    capability_inserts = []
    calibration_inserts = []

    equipment_no = None
    equipment_seq = 0  # 1-based index matching AUTO_INCREMENT order
    pending_calibrations = []  # dedupe (resp, freq, last, next) within one equipment group

    def flush_calibrations():
        seen = set()
        for eq, resp, freq, last, next_ in pending_calibrations:
            key = (resp, freq, last, next_)
            if key in seen:
                continue
            seen.add(key)
            calibration_inserts.append(
                "({eq}, {resp}, {freq}, {last}, {next_})".format(
                    eq=eq, resp=esc(resp), freq=esc(freq), last=esc_date(last), next_=esc_date(next_)
                )
            )
        pending_calibrations.clear()

    for r in range(5, ws.max_row + 1):
        no = ws.cell(row=r, column=1).value
        name = ws.cell(row=r, column=2).value
        model = ws.cell(row=r, column=3).value
        serial = ws.cell(row=r, column=4).value
        purpose = ws.cell(row=r, column=5).value
        meas_range = ws.cell(row=r, column=6).value
        accuracy = ws.cell(row=r, column=7).value
        operating_range = ws.cell(row=r, column=8).value
        acceptance = ws.cell(row=r, column=9).value
        responsible = ws.cell(row=r, column=10).value
        frequency = ws.cell(row=r, column=11).value
        last_cal = ws.cell(row=r, column=12).value
        next_cal = ws.cell(row=r, column=13).value
        location = ws.cell(row=r, column=14).value
        status_raw = ws.cell(row=r, column=15).value
        remarks = ws.cell(row=r, column=16).value

        if all(v is None for v in [no, name, model, serial, purpose]):
            continue

        # A new equipment "header" row starts a group (has a register number).
        if no is not None:
            flush_calibrations()
            equipment_seq += 1
            status, status_note = normalize_status(status_raw)
            equipment_inserts.append(
                "({no}, {name}, {model}, NULL, {serial}, NULL, {location}, "
                "{status}, {status_note}, {remarks})".format(
                    no=esc(no),
                    name=esc(name),
                    model=esc(model),
                    serial=esc(serial),
                    location=esc(location),
                    status=esc(status),
                    status_note=esc(status_note),
                    remarks=esc(remarks),
                )
            )
            equipment_no = equipment_seq

        if equipment_no is None:
            continue  # capability row before any header row -- shouldn't happen

        if purpose is not None:
            capability_inserts.append(
                "({eq}, {purpose}, {rng}, {acc}, {op_rng}, {accept})".format(
                    eq=equipment_no,
                    purpose=esc(purpose),
                    rng=esc(meas_range),
                    acc=esc(accuracy),
                    op_rng=esc(operating_range),
                    accept=esc(acceptance),
                )
            )

        if last_cal is not None or next_cal is not None:
            pending_calibrations.append(
                (equipment_no, responsible, frequency, last_cal, next_cal)
            )

    flush_calibrations()

    with open(OUTPUT_SQL, "w", encoding="utf-8") as f:
        f.write("-- Auto-generated from the equipment calibration register.\n")
        f.write("-- Regenerate with: python3 generate_seed_data.py\n\n")
        f.write("USE mcems;\n\n")

        f.write(
            "INSERT INTO equipment "
            "(equipment_no, name, model, manufacturer, serial_number, category, "
            "location, status, status_note, remarks) VALUES\n"
        )
        f.write(",\n".join(equipment_inserts) + ";\n\n")

        f.write(
            "INSERT INTO equipment_capability "
            "(equipment_id, measurement_purpose, measurement_range, "
            "instrument_accuracy, product_operating_range, acceptance_criteria) VALUES\n"
        )
        f.write(",\n".join(capability_inserts) + ";\n\n")

        f.write(
            "INSERT INTO calibration_history "
            "(equipment_id, responsible_person, calibration_frequency, "
            "calibration_date, next_calibration_date) VALUES\n"
        )
        f.write(",\n".join(calibration_inserts) + ";\n")

    print(f"equipment rows:            {len(equipment_inserts)}")
    print(f"equipment_capability rows: {len(capability_inserts)}")
    print(f"calibration_history rows:  {len(calibration_inserts)}")
    print(f"wrote {OUTPUT_SQL}")


if __name__ == "__main__":
    main()
